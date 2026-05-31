#!/usr/bin/env python3
"""
Generate a one-page executive summary workbook for portfolio presentation.

Produces a business-facing French summary of B2B data qualification KPIs
for electronic invoicing preparation.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

# --- Paths ---
PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_FILE = PROJECT_ROOT / "outputs" / "executive_summary.xlsx"

SHEET_NAME = "Synthèse projet"
TITLE = "Synthèse — Qualification de données clients B2B"
SUBTITLE = "Préparation à la facturation électronique"
DISCLAIMER = (
    "Les données utilisées sont synthétiques et ne contiennent "
    "aucune donnée réelle d'entreprise."
)

# Styling
FILL_HEADER = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
FILL_TITLE_BG = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
FONT_TITLE = Font(bold=True, size=16, color="1F3864")
FONT_SUBTITLE = Font(italic=True, size=11, color="44546A")
FONT_HEADER = Font(bold=True, color="1F3864")
FONT_SECTION = Font(bold=True, size=11, color="1F3864")
FONT_KPI = Font(bold=True, size=11)
FONT_NOTE = Font(italic=True, size=10, color="595959")
ALIGN_TOP = Alignment(vertical="top")

# Main KPI table (Indicateur | Nombre)
KPI_ROWS = [
    ("Comptes clients analysés", 1000),
    ("Comptes prêts pour mise à jour ERP", 70),
    ("Comptes nécessitant une correction", 739),
    ("Comptes nécessitant une relance téléphonique", 93),
    ("Doublons potentiels à vérifier", 98),
    ("Comptes préparés pour relance", 261),
    ("Relances prioritaires élevées", 125),
    ("E-mails de facturation manquants", 72),
    ("Numéros de TVA manquants", 70),
    ("SIRET manquants", 76),
]

# Duplicate detection improvement section
IMPROVEMENT_TITLE = "Amélioration de la détection des doublons"
IMPROVEMENT_ROWS = [
    ("Alertes initiales", 604),
    ("Alertes après affinage", 113),
    ("Réduction", "81,3 %"),
]


def write_executive_summary(path: Path) -> None:
    """Build the executive summary workbook with titles, KPIs and formatting."""
    path.parent.mkdir(parents=True, exist_ok=True)

    # Write base content starting at row 4 (room for title block)
    kpi_df = pd.DataFrame(KPI_ROWS, columns=["Indicateur", "Nombre"])

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        kpi_df.to_excel(writer, index=False, sheet_name=SHEET_NAME, startrow=3)

    wb = load_workbook(path)
    ws = wb[SHEET_NAME]

    # --- Title block (rows 1–2) ---
    ws.merge_cells("A1:B1")
    ws["A1"] = TITLE
    ws["A1"].font = FONT_TITLE
    ws["A1"].fill = FILL_TITLE_BG
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("A2:B2")
    ws["A2"] = SUBTITLE
    ws["A2"].font = FONT_SUBTITLE
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 20

    # --- KPI table header (row 4) ---
    for col in (1, 2):
        cell = ws.cell(row=4, column=col)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(horizontal="center" if col == 2 else "left", vertical="center")
    ws.row_dimensions[4].height = 22

    # --- KPI data rows (rows 5–14) ---
    first_data_row = 5
    last_kpi_row = first_data_row + len(KPI_ROWS) - 1
    for row in range(first_data_row, last_kpi_row + 1):
        ws.cell(row=row, column=1).alignment = Alignment(vertical="top", wrap_text=True)
        value_cell = ws.cell(row=row, column=2)
        value_cell.font = FONT_KPI
        value_cell.alignment = Alignment(horizontal="right", vertical="top")
        ws.row_dimensions[row].height = 18

    # --- Improvement section ---
    section_row = last_kpi_row + 2
    ws.merge_cells(f"A{section_row}:B{section_row}")
    ws.cell(row=section_row, column=1, value=IMPROVEMENT_TITLE).font = FONT_SECTION

    improvement_start = section_row + 1
    for i, (label, value) in enumerate(IMPROVEMENT_ROWS):
        r = improvement_start + i
        ws.cell(row=r, column=1, value=label).alignment = ALIGN_TOP
        val_cell = ws.cell(row=r, column=2, value=value)
        val_cell.font = FONT_KPI
        val_cell.alignment = Alignment(horizontal="right", vertical="top")
        ws.row_dimensions[r].height = 18

    # --- Disclaimer note ---
    note_row = improvement_start + len(IMPROVEMENT_ROWS) + 1
    ws.merge_cells(f"A{note_row}:B{note_row}")
    note_cell = ws.cell(row=note_row, column=1, value=DISCLAIMER)
    note_cell.font = FONT_NOTE
    note_cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[note_row].height = 32

    # --- Column widths ---
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 16

    # Freeze below header row of KPI table
    ws.freeze_panes = "A5"

    wb.save(path)


def main() -> None:
    write_executive_summary(OUTPUT_FILE)

    print("Synthèse exécutive générée.")
    print(f"  Fichier : {OUTPUT_FILE}")
    print(f"  Feuille : {SHEET_NAME}")
    print(f"  Indicateurs : {len(KPI_ROWS)}")


if __name__ == "__main__":
    main()
