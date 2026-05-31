#!/usr/bin/env python3
"""
Generate an operational phone follow-up file for B2B electronic invoicing.

Reads qualified client data and builds a call campaign workbook for accounting /
administrative outreach to collect missing fiscal and billing information.
Business-facing labels are in French.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# --- Paths ---
PROJECT_ROOT = Path(__file__).resolve().parent.parent
INPUT_FILE = PROJECT_ROOT / "outputs" / "b2b_clients_qualified.xlsx"
OUTPUT_FILE = PROJECT_ROOT / "outputs" / "accounts_to_call.xlsx"

# --- Selection criteria (internal keys from script 02 — input stays in English) ---
STATUS_PHONE_FOLLOWUP = "Needs phone follow-up"
ACTION_CALL = "Call billing/admin contact"

ISSUE_MISSING_EMAIL = "missing_billing_email"
ISSUE_MISSING_CONTACT = "missing_billing_contact"
ISSUE_MISSING_VAT = "missing_vat_number"
ISSUE_MISSING_SIRET = "missing_siret"
ISSUE_INVALID_VAT = "invalid_vat_number"
ISSUE_INVALID_SIRET = "invalid_siret"
ISSUE_INVALID_EMAIL = "invalid_billing_email"
ISSUE_MISSING_PHONE = "missing_phone"

# Priority: English from qualified file → French in output
PRIORITY_FROM_EN = {"High": "Élevée", "Medium": "Moyenne", "Low": "Faible"}
PRIORITY_ORDER_EN = {"High": 0, "Medium": 1, "Low": 2}

# French labels for anomalies displayed in the export
ISSUE_LABELS_FR = {
    "missing_siren": "SIREN manquant",
    "invalid_siren": "SIREN invalide",
    "missing_siret": "SIRET manquant",
    "invalid_siret": "SIRET invalide",
    "missing_vat_number": "TVA manquante",
    "invalid_vat_number": "TVA invalide",
    "missing_billing_email": "E-mail de facturation manquant",
    "invalid_billing_email": "E-mail de facturation invalide",
    "missing_billing_contact": "Contact facturation manquant",
    "missing_phone": "Téléphone manquant",
    "formatting_company_name": "Format du nom client incorrect",
    "duplicate_siret": "Doublon SIRET",
    "duplicate_siren_and_name": "Doublon SIREN et raison sociale",
}

# Excel styling — French sheet names
SHEET_SUMMARY = "Synthèse relance"
SHEET_ACCOUNTS = "Comptes à relancer"

HEADER_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
HEADER_FONT = Font(bold=True, color="1F3864")
PRIORITY_FILLS = {
    "Élevée": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "Moyenne": PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid"),
    "Faible": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
}
ALIGN_TOP = Alignment(vertical="top")
ALIGN_TOP_WRAP = Alignment(vertical="top", wrap_text=True)
ALIGN_HEADER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Column index on "Comptes à relancer" sheet (1-based)
COL_PRIORITY = 7
COL_CALL_SCRIPT = 9
DATA_ROW_HEIGHT = 72
HEADER_ROW_HEIGHT = 24

# French output columns for the detailed call list
CAMPAIGN_COLUMNS = [
    "Identifiant client",
    "Nom du client",
    "Contact facturation",
    "Téléphone",
    "E-mail facturation",
    "Anomalies détectées",
    "Priorité",
    "Objectif de l'appel",
    "Script d'appel",
    "Prochaine action",
    "Statut de l'appel",
    "Date de relance",
    "Notes",
]

# French call script templates (professional, adapted per objective)
SCRIPT_INTRO = (
    "Bonjour, je vous contacte dans le cadre de la mise à jour de vos informations "
    "de facturation. Afin de préparer la transition vers la facturation électronique, "
    "nous souhaitons confirmer certaines informations administratives de votre compte."
)

SCRIPT_BY_TOPIC = {
    "vat": " Pourriez-vous me confirmer votre numéro de TVA intracommunautaire (format FR) ?",
    "siret": " Pourriez-vous me communiquer le SIRET de l'établissement facturé (14 chiffres) ?",
    "email": " Pourriez-vous me confirmer l'adresse e-mail de facturation à utiliser pour vos factures ?",
    "contact": (
        " Pourriez-vous m'indiquer le nom et les coordonnées du contact facturation "
        "(service comptabilité ou administration) ?"
    ),
    "phone": (
        " Nous ne disposons pas de numéro de téléphone à jour : pourriez-vous nous indiquer "
        "un contact téléphonique ou un canal de communication préféré ?"
    ),
}

# Prochaine action (French business labels)
NEXT_ACTION_CALL = "Appeler le service comptable/administratif"
NEXT_ACTION_EMAIL = "Envoyer un e-mail de relance"
NEXT_ACTION_SEARCH = "Rechercher un contact alternatif"
NEXT_ACTION_REVIEW = "Vérifier en interne avant contact"

DEFAULT_CALL_STATUS = "À appeler"


def clean_text(value) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip()


def parse_issues(data_issues: str) -> set[str]:
    """Split semicolon-separated issue tags into a set."""
    if not data_issues:
        return set()
    return {part.strip() for part in data_issues.split(";") if part.strip()}


def translate_issues_display(data_issues: str) -> str:
    """Convert internal issue tags to French labels for the export."""
    tags = parse_issues(data_issues)
    if not tags:
        return ""
    return "; ".join(ISSUE_LABELS_FR.get(tag, tag) for tag in sorted(tags))


def translate_priority(priority_en: str) -> str:
    """Map English priority from qualified file to French export label."""
    return PRIORITY_FROM_EN.get(clean_text(priority_en), "Moyenne")


def requires_client_contact(row: pd.Series) -> bool:
    """Return True if the account should appear in the call campaign."""
    issues = parse_issues(row.get("data_issues", ""))

    if row.get("qualification_status") == STATUS_PHONE_FOLLOWUP:
        return True
    if row.get("action_to_take") == ACTION_CALL:
        return True
    if ISSUE_MISSING_EMAIL in issues:
        return True
    if ISSUE_MISSING_CONTACT in issues:
        return True
    if ISSUE_MISSING_VAT in issues:
        return True
    if ISSUE_MISSING_SIRET in issues:
        return True

    return False


def build_call_objectives(issues: set[str]) -> str:
    """Build a French call_objective from detected data_issues."""
    parts: list[str] = []

    if ISSUE_MISSING_VAT in issues or ISSUE_INVALID_VAT in issues:
        parts.append("Collecter ou confirmer le numéro de TVA")
    if ISSUE_MISSING_SIRET in issues or ISSUE_INVALID_SIRET in issues:
        parts.append("Collecter ou confirmer le SIRET")
    if ISSUE_MISSING_EMAIL in issues or ISSUE_INVALID_EMAIL in issues:
        parts.append("Confirmer l'e-mail de facturation")
    if ISSUE_MISSING_CONTACT in issues:
        parts.append("Identifier le contact facturation/administratif")
    if ISSUE_MISSING_PHONE in issues:
        parts.append("Trouver un canal de contact alternatif")

    if not parts:
        return (
            "Confirmer les informations de facturation et fiscales "
            "pour la facturation électronique"
        )

    if len(parts) == 1:
        return parts[0]

    return "; ".join(parts[:-1]) + "; " + parts[-1]


def build_call_script(issues: set[str]) -> str:
    """Compose a short professional French script adapted to the call objective."""
    topics: list[str] = []

    if ISSUE_MISSING_VAT in issues or ISSUE_INVALID_VAT in issues:
        topics.append("vat")
    if ISSUE_MISSING_SIRET in issues or ISSUE_INVALID_SIRET in issues:
        topics.append("siret")
    if ISSUE_MISSING_EMAIL in issues or ISSUE_INVALID_EMAIL in issues:
        topics.append("email")
    if ISSUE_MISSING_CONTACT in issues:
        topics.append("contact")
    if ISSUE_MISSING_PHONE in issues:
        topics.append("phone")

    if not topics:
        return (
            f"{SCRIPT_INTRO} Pourriez-vous m'aider à confirmer les informations "
            "de facturation de votre compte (SIRET, TVA, e-mail de facturation) ?"
        )

    body = "".join(SCRIPT_BY_TOPIC[t] for t in topics)
    return SCRIPT_INTRO + body


def determine_next_action(row: pd.Series, issues: set[str]) -> str:
    """Suggest the next operational step (French label)."""
    phone = clean_text(row.get("phone", ""))
    email = clean_text(row.get("billing_email", ""))

    if ISSUE_MISSING_PHONE in issues or not phone:
        if email and is_valid_enough_email(email):
            return NEXT_ACTION_EMAIL
        return NEXT_ACTION_SEARCH

    if (ISSUE_INVALID_EMAIL in issues or ISSUE_MISSING_EMAIL in issues) and phone:
        return NEXT_ACTION_CALL

    fiscal = {
        ISSUE_MISSING_VAT, ISSUE_INVALID_VAT,
        ISSUE_MISSING_SIRET, ISSUE_INVALID_SIRET,
    }
    if fiscal & issues and phone:
        return NEXT_ACTION_CALL

    if ISSUE_MISSING_CONTACT in issues and phone:
        return NEXT_ACTION_CALL

    notes = clean_text(row.get("notes", "")).lower()
    if "doublon" in notes or "migration" in notes:
        return NEXT_ACTION_REVIEW

    return NEXT_ACTION_CALL


def is_valid_enough_email(email: str) -> bool:
    """Minimal check to decide if a follow-up email is plausible."""
    return "@" in email and "." in email.split("@")[-1]


def select_accounts_to_call(df: pd.DataFrame) -> pd.DataFrame:
    """Filter qualified records that require client contact."""
    mask = df.apply(requires_client_contact, axis=1)
    return df.loc[mask].copy()


def build_campaign_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """Build the call campaign table with French business labels."""
    rows = []

    for _, row in df.iterrows():
        issues = parse_issues(row.get("data_issues", ""))
        priority_en = clean_text(row.get("priority", "Medium"))

        rows.append({
            "Identifiant client": row.get("client_id", ""),
            "Nom du client": row.get("company_name", ""),
            "Contact facturation": row.get("billing_contact_name", ""),
            "Téléphone": row.get("phone", ""),
            "E-mail facturation": row.get("billing_email", ""),
            "Anomalies détectées": translate_issues_display(row.get("data_issues", "")),
            "Priorité": translate_priority(priority_en),
            "Objectif de l'appel": build_call_objectives(issues),
            "Script d'appel": build_call_script(issues),
            "Prochaine action": determine_next_action(row, issues),
            "Statut de l'appel": DEFAULT_CALL_STATUS,
            "Date de relance": "",
            "Notes": row.get("notes", ""),
            "_priority_rank": PRIORITY_ORDER_EN.get(priority_en, 99),
        })

    campaign = pd.DataFrame(rows)
    campaign = campaign.sort_values(
        ["_priority_rank", "Nom du client"],
        ascending=[True, True],
    ).drop(columns="_priority_rank")
    campaign = campaign[CAMPAIGN_COLUMNS]

    return campaign.reset_index(drop=True)


def count_issue(campaign: pd.DataFrame, issue_tag: str) -> int:
    """Count rows whose French anomalies label contains the given issue."""
    label = ISSUE_LABELS_FR.get(issue_tag, issue_tag)
    return campaign["Anomalies détectées"].str.contains(label, regex=False).sum()


def build_summary_dataframe(campaign: pd.DataFrame) -> pd.DataFrame:
    """KPI table for the Synthèse relance sheet (French labels)."""
    rows = [
        ("Comptes à relancer", len(campaign)),
        ("Relances prioritaires élevées", (campaign["Priorité"] == "Élevée").sum()),
        ("Relances priorité moyenne", (campaign["Priorité"] == "Moyenne").sum()),
        ("Relances priorité faible", (campaign["Priorité"] == "Faible").sum()),
        ("Téléphone manquant", count_issue(campaign, ISSUE_MISSING_PHONE)),
        ("E-mail de facturation manquant", count_issue(campaign, ISSUE_MISSING_EMAIL)),
        ("TVA manquante", count_issue(campaign, ISSUE_MISSING_VAT)),
        ("SIRET manquant", count_issue(campaign, ISSUE_MISSING_SIRET)),
    ]
    return pd.DataFrame(rows, columns=["Indicateur", "Nombre"])


def style_summary_sheet(ws) -> None:
    """Format the Synthèse relance sheet."""
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 14

    for col in range(1, 3):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_HEADER

    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=1).alignment = ALIGN_TOP
        count_cell = ws.cell(row=row, column=2)
        count_cell.alignment = Alignment(horizontal="right", vertical="top")
        count_cell.font = Font(bold=True)

    ws.freeze_panes = "A2"


def style_accounts_sheet(ws, n_rows: int, n_cols: int) -> None:
    """
    Format Comptes à relancer: light blue headers, priority colours,
    wrapped script, taller rows, filter, frozen header.
    """
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_HEADER
    ws.row_dimensions[1].height = HEADER_ROW_HEIGHT

    for row in range(2, n_rows + 2):
        ws.row_dimensions[row].height = DATA_ROW_HEIGHT

        for col in range(1, n_cols + 1):
            cell = ws.cell(row=row, column=col)
            if col in (COL_CALL_SCRIPT, 8, 6):  # Script, objectif, anomalies
                cell.alignment = ALIGN_TOP_WRAP
            else:
                cell.alignment = ALIGN_TOP

        priority_val = clean_text(ws.cell(row=row, column=COL_PRIORITY).value)
        priority_cell = ws.cell(row=row, column=COL_PRIORITY)
        priority_cell.alignment = Alignment(horizontal="center", vertical="top")
        if priority_val in PRIORITY_FILLS:
            priority_cell.fill = PRIORITY_FILLS[priority_val]
            priority_cell.font = Font(bold=True)

        ws.cell(row=row, column=COL_CALL_SCRIPT).alignment = ALIGN_TOP_WRAP

    widths = {
        "A": 16, "B": 28, "C": 22, "D": 18, "E": 30, "F": 38,
        "G": 12, "H": 44, "I": 60, "J": 38, "K": 14, "L": 14, "M": 24,
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{n_rows + 1}"


def export_campaign(campaign: pd.DataFrame, summary: pd.DataFrame, path: Path) -> None:
    """Write two-sheet workbook and apply formatting."""
    path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary.to_excel(writer, index=False, sheet_name=SHEET_SUMMARY)
        campaign.to_excel(writer, index=False, sheet_name=SHEET_ACCOUNTS)

    from openpyxl import load_workbook

    wb = load_workbook(path)
    style_summary_sheet(wb[SHEET_SUMMARY])
    style_accounts_sheet(
        wb[SHEET_ACCOUNTS],
        n_rows=len(campaign),
        n_cols=len(CAMPAIGN_COLUMNS),
    )
    wb.save(path)


def main() -> None:
    print(f"Lecture : {INPUT_FILE}")
    qualified = pd.read_excel(INPUT_FILE, dtype=str).fillna("")

    selected = select_accounts_to_call(qualified)
    campaign = build_campaign_dataframe(selected)
    summary = build_summary_dataframe(campaign)

    export_campaign(campaign, summary, OUTPUT_FILE)

    print("\n--- Synthèse relance ---")
    print(f"Comptes à relancer :              {len(campaign)}")
    print(f"Priorité élevée :                 {(campaign['Priorité'] == 'Élevée').sum()}")
    print(f"Priorité moyenne :                {(campaign['Priorité'] == 'Moyenne').sum()}")
    print(f"Priorité faible :                 {(campaign['Priorité'] == 'Faible').sum()}")
    print(f"Téléphone manquant :              {count_issue(campaign, ISSUE_MISSING_PHONE)}")
    print(f"E-mail de facturation manquant :  {count_issue(campaign, ISSUE_MISSING_EMAIL)}")
    print(f"TVA manquante :                   {count_issue(campaign, ISSUE_MISSING_VAT)}")
    print(f"SIRET manquant :                  {count_issue(campaign, ISSUE_MISSING_SIRET)}")
    print(f"\nFichier généré :\n  {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
